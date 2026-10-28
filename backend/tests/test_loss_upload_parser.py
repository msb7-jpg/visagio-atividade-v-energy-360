import io
import re
import zipfile
from datetime import date

import pytest
from openpyxl import Workbook

from app.features.loss_management.schemas import InspectionStatus
from app.features.loss_management.service import LossService
from app.features.loss_management.upload_parser import MAX_DATA_ROWS, read_spreadsheet
from app.features.loss_management.upload_reconcile import reconcile_rows
from app.shared.exceptions import BadRequestException

CSV_HEADER = "id_inspecao;data_inspecao;irregularidade_constatada;tipo_irregularidade;kwh_recuperado;numero_toi"


def csv_bytes(*lines: str, encoding: str = "utf-8-sig") -> bytes:
    return "\r\n".join([CSV_HEADER, *lines]).encode(encoding)


def xlsx_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, sheet_rows in sheets.items():
        worksheet = workbook.create_sheet(sheet_name)
        for sheet_row in sheet_rows:
            worksheet.append(sheet_row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_reads_csv_exported_by_the_portal_with_bom_and_semicolon():
    content = csv_bytes("7;14.08.2025;CONFIRMADA;by-pass;1.234,5;TOI-1")

    rows = read_spreadsheet(content, "retorno.csv")

    assert len(rows) == 1
    assert rows[0].cells["id_inspecao"] == "7"
    assert rows[0].cells["kwh_recuperado"] == "1.234,5"
    assert rows[0].line == 2


def test_reads_csv_with_comma_delimiter_and_latin1_encoding():
    content = "ID Inspeção,Irregularidade Constatada\n9,NÃO".encode("cp1252")

    rows = read_spreadsheet(content, "retorno.CSV")

    assert rows[0].cells == {"id_inspecao": "9", "irregularidade_constatada": "NÃO"}


def test_reads_every_sheet_of_a_workbook_and_skips_sheets_without_the_key_column():
    content = xlsx_bytes({
        "Equipe 1": [["id_inspecao", "kWh Recuperado"], [1, 10.5]],
        "Equipe 2": [["ID Inspeção", "kWh Recuperado"], [2, "7,5"], [None, None], [3, 1]],
        "Resumo": [["total"], [3]],
    })

    rows = read_spreadsheet(content, "retorno.xlsx")

    assert [(row.sheet, row.line, row.cells["id_inspecao"]) for row in rows] == [
        ("Equipe 1", 2, 1),
        ("Equipe 2", 2, 2),
        ("Equipe 2", 4, 3),
    ]


@pytest.mark.parametrize("filename", ["retorno.pdf", "retorno", "retorno.xls"])
def test_rejects_unsupported_file_types(filename):
    with pytest.raises(BadRequestException, match="xlsx"):
        read_spreadsheet(b"x", filename)


def test_rejects_file_without_the_key_column():
    with pytest.raises(BadRequestException, match="id_inspecao"):
        read_spreadsheet(b"a;b\r\n1;2", "retorno.csv")


def test_rejects_corrupted_workbook():
    with pytest.raises(BadRequestException, match="xlsx"):
        read_spreadsheet(b"not a zip", "retorno.xlsx")


FIELD_HEADER = "id_inspecao;id_uc;data;inspetor;resultado;tipo_irregularidade;kwh_recuperado"


def field_csv(*lines: str) -> bytes:
    return "\r\n".join([FIELD_HEADER, *lines]).encode("utf-8-sig")


def test_field_sheet_column_names_are_read_as_the_portal_columns():
    rows = read_spreadsheet(field_csv("IN0000007;UC1;14.08.2025;INSP-028;CONFIRMADA;BY-PASS;838,10"), "campo.csv")

    assert set(rows[0].cells) == {
        "id_inspecao",
        "id_uc",
        "data_inspecao",
        "inspetor",
        "irregularidade_constatada",
        "tipo_irregularidade",
        "kwh_recuperado",
    }


def test_alias_never_overrides_a_column_that_is_already_in_the_file():
    rows = read_spreadsheet(b"id_inspecao;data;data_inspecao\r\n1;2025-01-01;2025-02-02", "campo.csv")

    assert rows[0].cells["data_inspecao"] == "2025-02-02"
    assert rows[0].cells["data"] == "2025-01-01"


def test_file_keyed_only_by_unit_is_accepted():
    rows = read_spreadsheet(b"id_uc;resultado\r\nUC1;sim", "campo.csv")

    assert rows[0].cells == {"id_uc": "UC1", "irregularidade_constatada": "sim"}


def test_sheet_with_unit_but_without_result_column_is_not_a_data_sheet():
    content = xlsx_bytes({
        "Equipe": [["id_uc", "resultado"], ["UC1", "sim"]],
        "Resumo": [["id_uc", "total"], ["UC1", 3], ["UC2", 5]],
    })

    rows = read_spreadsheet(content, "campo.xlsx")

    assert [(row.sheet, row.cells["id_uc"]) for row in rows] == [("Equipe", "UC1")]


def test_workbook_where_only_summary_sheets_have_unit_column_is_rejected():
    content = xlsx_bytes({"Resumo": [["id_uc", "total"], ["UC1", 3]]})

    with pytest.raises(BadRequestException, match="id_inspecao"):
        read_spreadsheet(content, "campo.xlsx")


def test_csv_with_unit_but_without_result_column_is_rejected():
    with pytest.raises(BadRequestException, match="id_inspecao"):
        read_spreadsheet(b"id_uc;total\r\nUC1;3", "campo.csv")


def test_unit_sheet_with_result_header_and_no_rows_is_accepted_as_empty():
    content = xlsx_bytes({"Equipe": [["id_uc", "resultado"]]})

    assert read_spreadsheet(content, "campo.xlsx") == []


def reconcile_field(*lines: str, active_by_uc: dict[str, int] | None = None, statuses: dict | None = None):
    rows = read_spreadsheet(field_csv(*lines), "campo.csv")
    order_statuses = statuses or dict.fromkeys((active_by_uc or {}).values(), InspectionStatus.PENDENTE)
    return reconcile_rows(rows, order_statuses, active_by_uc or {})


def test_external_id_is_matched_to_the_active_order_of_the_unit():
    plan = reconcile_field(
        "IN0000007;UC1;09-25-2025;INSP-028  ;IRREGULARIDADE CONFIRMADA;BY-PASS;838,10", active_by_uc={"UC1": 5}
    )

    update = plan.updates[5]
    assert update.status == InspectionStatus.CONCLUIDA
    assert update.inspetor == "INSP-028"
    assert update.data_inspecao == date(2025, 9, 25)
    assert update.fiscal_responsavel is None
    assert update.irregularidade_constatada is True
    assert update.tipo_irregularidade == "BY-PASS"
    assert update.kwh_recuperado == pytest.approx(838.1)
    assert plan.rejections == []


def test_negative_field_row_discards_placeholder_details():
    plan = reconcile_field("IN0000008;UC2;2025-11-08;INSP-031;SEM IRREGULARIDADE;NULL;0,00", active_by_uc={"UC2": 6})

    update = plan.updates[6]
    assert update.irregularidade_constatada is False
    assert update.tipo_irregularidade is None
    assert update.kwh_recuperado is None


def test_unit_without_active_order_is_rejected_with_reason():
    plan = reconcile_field("IN0000009;UC9;2025-11-08;INSP-031;CONFIRMADA;BY-PASS;10,0", active_by_uc={"UC1": 5})

    assert plan.updates == {}
    assert "UC9" in plan.rejections[0].motivo
    assert "ordem ativa" in plan.rejections[0].motivo


def test_numeric_portal_id_wins_over_the_unit_column():
    plan = reconcile_field(
        "7;UC1;2025-11-08;INSP-031;CONFIRMADA;BY-PASS;10,0",
        active_by_uc={"UC1": 5},
        statuses={5: InspectionStatus.PENDENTE, 7: InspectionStatus.PENDENTE},
    )

    assert set(plan.updates) == {7}


def test_two_rows_for_the_same_unit_keep_the_first():
    plan = reconcile_field(
        "IN1;UC1;2025-11-08;INSP-031;CONFIRMADA;BY-PASS;10,0",
        "IN2;UC1;2025-11-09;INSP-032;SEM IRREGULARIDADE;NULL;0,00",
        active_by_uc={"UC1": 5},
    )

    assert plan.updates[5].irregularidade_constatada is True
    assert "primeira ocorrência" in plan.rejections[0].motivo


def test_field_row_without_result_is_ignored_when_the_unit_has_an_order():
    plan = reconcile_field("IN1;UC1;;INSP-031;;;", active_by_uc={"UC1": 5})

    assert plan.updates == {}
    assert plan.ignored == 1


def reconcile(*lines: str, active_ids: tuple[int, ...] = (1, 2, 3), statuses: dict | None = None):
    orders = statuses or dict.fromkeys(active_ids, InspectionStatus.PENDENTE)
    return reconcile_rows(read_spreadsheet(csv_bytes(*lines), "retorno.csv"), orders)


def test_confirmed_irregularity_becomes_a_completed_update():
    plan = reconcile("1;2025-08-01;irregularidade confirmada;By-pass;1.200,5;TOI-9")

    update = plan.updates[1]
    assert update.status == InspectionStatus.CONCLUIDA
    assert update.irregularidade_constatada is True
    assert update.kwh_recuperado == pytest.approx(1200.5)
    assert update.tipo_irregularidade == "By-pass"
    assert update.numero_toi == "TOI-9"
    assert plan.rejections == []


def test_generated_orders_can_also_be_completed_from_the_field_sheet():
    plan = reconcile("1;;NADA CONSTA;;;", statuses={1: InspectionStatus.GERADA})

    assert plan.updates[1].irregularidade_constatada is False


def test_blank_result_rows_are_ignored_not_rejected():
    plan = reconcile("1;;;;;", "2;;;;;")

    assert plan.updates == {}
    assert plan.rejections == []
    assert plan.ignored == 2


def test_invalid_rows_are_rejected_with_reason_without_stopping_the_batch():
    plan = reconcile(
        "1;;sim;;abc;",
        "999;;sim;;;",
        "abc;;sim;;;",
        "2;;talvez;;;",
        "3;31.02.2025;sim;;;",
        "1;;sim;;50;",
    )

    assert set(plan.updates) == {1}
    reasons = {rejection.linha: rejection.motivo for rejection in plan.rejections}
    assert "kwh_recuperado" in reasons[2]
    assert "não encontrada" in reasons[3]
    assert "id_inspecao" in reasons[4]
    assert "irregularidade_constatada" in reasons[5]
    assert "data_inspecao" in reasons[6]
    assert plan.rejections[0].id_inspecao == 1


def test_duplicate_id_in_the_same_file_is_rejected():
    plan = reconcile("1;;sim;;10;", "1;;nao;;;")

    assert plan.updates[1].irregularidade_constatada is True
    assert "duplicada" in plan.rejections[0].motivo


@pytest.mark.parametrize("status", [InspectionStatus.CONCLUIDA, InspectionStatus.CANCELADA])
def test_orders_that_already_left_the_field_flow_are_rejected(status):
    plan = reconcile("1;;sim;;10;", statuses={1: status})

    assert plan.updates == {}
    assert status.value in plan.rejections[0].motivo


def test_recovered_energy_without_irregularity_is_contradictory():
    plan = reconcile("1;;NADA CONSTA;;300;")

    assert plan.updates == {}
    assert "kwh_recuperado" in plan.rejections[0].motivo


@pytest.mark.parametrize("raw_id", ["1.0", "1,0", " 1 "])
def test_order_id_tolerates_decimal_text_written_by_other_tools(raw_id):
    plan = reconcile(f"{raw_id};;sim;;10;")

    assert set(plan.updates) == {1}


def test_negative_result_tolerates_zero_kwh_and_prefilled_details():
    plan = reconcile("1;;NADA CONSTA;by-pass;0,00;TOI-1")

    update = plan.updates[1]
    assert update.irregularidade_constatada is False
    assert update.kwh_recuperado is None
    assert update.tipo_irregularidade is None
    assert update.numero_toi is None


def test_duplicate_rejection_says_the_first_value_was_kept():
    plan = reconcile("1;;sim;;10;", "1;;nao;;;")

    assert "primeira ocorrência" in plan.rejections[0].motivo


def test_file_over_the_row_limit_is_rejected():
    lines = [f"{index};;sim;;1;" for index in range(1, MAX_DATA_ROWS + 2)]

    with pytest.raises(BadRequestException, match="linhas"):
        read_spreadsheet(csv_bytes(*lines), "retorno.csv")


def declare_wrong_dimension(xlsx_content: bytes) -> bytes:
    rewritten = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(xlsx_content)) as source, zipfile.ZipFile(rewritten, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="A1:A1"', data)
            target.writestr(entry, data)
    return rewritten.getvalue()


def test_sheet_with_wrong_declared_dimension_is_fully_read():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["id_inspecao", "kwh_recuperado"])
    worksheet.append([1, 5])
    worksheet.append([2, 6])
    saved = io.BytesIO()
    workbook.save(saved)

    rows = read_spreadsheet(declare_wrong_dimension(saved.getvalue()), "retorno.xlsx")

    assert [row.cells["id_inspecao"] for row in rows] == [1, 2]


class StubRepository:
    """Simula uma ordem que saiu do fluxo de campo entre a validação e a gravação."""

    async def get_active_order_ids_by_uc(self, unit_ids):
        return {}

    async def get_order_statuses(self, ids):
        return dict.fromkeys(ids, InspectionStatus.PENDENTE)

    async def apply_field_results(self, updates):
        return []


class UnitLookupStubRepository(StubRepository):
    """Simula uma UC com ordem ativa e grava tudo o que o serviço pedir."""

    def __init__(self) -> None:
        self.looked_up_units: set[str] = set()

    async def get_active_order_ids_by_uc(self, unit_ids):
        self.looked_up_units = set(unit_ids)
        return {"UC1": 5}

    async def apply_field_results(self, updates):
        return list(updates)


async def test_upload_looks_up_only_the_units_of_rows_without_a_portal_id():
    repository = UnitLookupStubRepository()
    service = LossService(repository=repository)
    content = field_csv(
        "IN0000007;UC1;2025-11-08;INSP-031;CONFIRMADA;BY-PASS;10,0",
        "9;UC2;2025-11-08;INSP-031;CONFIRMADA;BY-PASS;10,0",
    )

    result = await service.upload_field_results(content, "campo.csv")

    assert repository.looked_up_units == {"UC1"}
    assert result.conciliadas == 2
    assert result.kwh_recuperado_total == pytest.approx(20.0)


async def test_order_that_left_field_flow_during_upload_is_reported_as_rejected():
    service = LossService(repository=StubRepository())

    result = await service.upload_field_results(csv_bytes("1;;sim;;10;"), "retorno.csv")

    assert result.conciliadas == 0
    assert result.kwh_recuperado_total == pytest.approx(0)
    assert (result.rejeitadas[0].linha, result.rejeitadas[0].id_inspecao) == (2, 1)

